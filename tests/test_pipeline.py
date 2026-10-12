from pathlib import Path
import pandas as pd
import pytest
from openpyxl import load_workbook
from sales_report.clean import clean, parse_dates, parse_number
from sales_report.config import load_config
from sales_report.ingest import ingest_folder, standardise_columns
from sales_report.aggregate import summarise
from sales_report.cli import run

ROOT = Path(__file__).resolve().parents[1]
CFG = load_config(ROOT / "config.yaml")


@pytest.mark.parametrize("raw,expected", [("1.234,50", 1234.5), ("1,234.50", 1234.5), ("€ 12,5", 12.5),
                                          ("12.5", 12.5), ("7", 7.0)])
def test_parse_number(raw, expected):
    assert parse_number(pd.Series([raw]))[0] == pytest.approx(expected)


def test_parse_number_invalid():
    assert parse_number(pd.Series(["abc"])).isna().all()


def test_parse_dates_formats():
    out = parse_dates(pd.Series(["2026-03-05", "05/03/2026", "31/02/2026"]))
    assert out[0] == out[1] == pd.Timestamp(2026, 3, 5)
    assert pd.isna(out[2])


def test_standardise_columns():
    df = pd.DataFrame(columns=["Order Date", "AREA", "qty", "Product Name", "price"])
    assert set(standardise_columns(df, CFG["columns"]).columns) == set(CFG["required"])


def _raw(**over):
    base = dict(date=["2026-01-01"], region=[" north "], product=["mouse"], quantity=["2"], unit_price=["10"])
    base.update(over); return pd.DataFrame(base)


def test_clean_valid_row():
    ok, bad = clean(_raw(), CFG)
    assert len(ok) == 1 and bad.empty
    assert ok.loc[0, "region"] == "North" and ok.loc[0, "revenue"] == 20


@pytest.mark.parametrize("over,reason", [({"quantity": ["-3"]}, "quantity out of range"),
                                         ({"date": ["31/02/2026"]}, "invalid date"),
                                         ({"unit_price": ["x"]}, "invalid price"),
                                         ({"region": [""]}, "missing region")])
def test_clean_rejects(over, reason):
    ok, bad = clean(_raw(**over), CFG)
    assert ok.empty and reason in bad.loc[0, "reason"]


def test_duplicates_removed():
    ok, _ = clean(pd.concat([_raw(), _raw()]), CFG)
    assert len(ok) == 1


def test_summary_totals():
    ok, _ = clean(pd.concat([_raw(), _raw(product=["laptop"], unit_price=["100"])]), CFG)
    t = summarise(ok)
    assert t["kpi"].set_index("Metric").loc["Total revenue", "Value"] == 220
    assert t["region"]["share"].sum() == pytest.approx(1)


def test_end_to_end(tmp_path):
    src = tmp_path / "in"; src.mkdir()
    _raw().to_csv(src / "a.csv", index=False)
    pd.DataFrame({"Date": ["2026-01-01"], "Notes": ["x"]}).to_csv(src / "broken.csv", index=False)
    out = run(src, tmp_path / "r.xlsx", ROOT / "config.yaml", tmp_path / "logs")
    wb = load_workbook(out)
    assert {"Summary", "Monthly", "Rejected Rows"} <= set(wb.sheetnames)
    assert len(wb["Summary"]._charts) == 3


def test_empty_folder_raises(tmp_path):
    with pytest.raises(FileNotFoundError):
        ingest_folder(tmp_path, CFG)
