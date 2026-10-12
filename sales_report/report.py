"""Write the formatted Excel report with native Excel charts."""
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER = PatternFill("solid", fgColor="1F4E78")


def _style(ws, money_cols=(), pct_cols=(), fmt="#,##0.00"):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = HEADER
        c.alignment = Alignment(horizontal="center")
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        ws.column_dimensions[letter].width = max(12, max(len(str(c.value or "")) for c in col) + 3)
        name = col[0].value
        for c in col[1:]:
            if name in money_cols: c.number_format = fmt
            if name in pct_cols: c.number_format = "0.0%"
    ws.freeze_panes = "A2"


def write_report(tables: dict, clean: pd.DataFrame, rejected: pd.DataFrame, path: Path, cfg: dict) -> Path:
    """Create the multi-sheet workbook and add charts."""
    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    money = cfg["report"]["currency_format"]
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        tables["kpi"].to_excel(xw, sheet_name="Summary", index=False, startrow=2)
        tables["monthly"].to_excel(xw, sheet_name="Monthly", index=False)
        tables["region"].to_excel(xw, sheet_name="By Region", index=False)
        tables["product"].to_excel(xw, sheet_name="Top Products", index=False)
        tables["pivot"].to_excel(xw, sheet_name="Month x Region", index=False)
        clean.assign(date=clean.date.dt.date).to_excel(xw, sheet_name="Clean Data", index=False)
        rejected.assign(date=rejected.date.astype(str)).to_excel(xw, sheet_name="Rejected Rows", index=False)
    wb = load_workbook(path)
    s = wb["Summary"]
    s["A1"] = cfg["report"]["title"]; s["A1"].font = Font(size=16, bold=True, color="1F4E78")
    for r in s.iter_rows(min_row=3, max_row=3):
        for c in r: c.font = Font(bold=True, color="FFFFFF"); c.fill = HEADER
    s.column_dimensions["A"].width = 24; s.column_dimensions["B"].width = 18
    for r in (4, 7): s.cell(r, 2).number_format = money
    for r in (5, 6): s.cell(r, 2).number_format = "#,##0"
    _style(wb["Monthly"], ["revenue"], fmt=money)
    _style(wb["By Region"], ["revenue"], ["share"], fmt=money)
    _style(wb["Top Products"], ["revenue"], fmt=money)
    _style(wb["Month x Region"], [c.value for c in wb["Month x Region"][1][1:]], fmt=money)
    _style(wb["Clean Data"], ["unit_price", "revenue"], fmt=money)
    _style(wb["Rejected Rows"])

    m = wb["Monthly"]; n = m.max_row
    lc = LineChart(); lc.title = "Monthly revenue"; lc.height, lc.width = 8, 16
    lc.add_data(Reference(m, min_col=2, min_row=1, max_row=n), titles_from_data=True)
    lc.set_categories(Reference(m, min_col=1, min_row=2, max_row=n)); s.add_chart(lc, "D3")
    rg = wb["By Region"]; n = rg.max_row
    pc = PieChart(); pc.title = "Revenue share by region"; pc.height, pc.width = 8, 12
    pc.add_data(Reference(rg, min_col=2, min_row=1, max_row=n), titles_from_data=True)
    pc.set_categories(Reference(rg, min_col=1, min_row=2, max_row=n)); s.add_chart(pc, "D20")
    tp = wb["Top Products"]; n = tp.max_row
    bc = BarChart(); bc.type = "bar"; bc.title = "Top products by revenue"; bc.height, bc.width = 8, 14
    bc.add_data(Reference(tp, min_col=2, min_row=1, max_row=n), titles_from_data=True)
    bc.set_categories(Reference(tp, min_col=1, min_row=2, max_row=n)); s.add_chart(bc, "M20")
    wb.save(path)
    return path
